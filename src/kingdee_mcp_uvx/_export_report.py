import os
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


MO_ACTIVE_STATUS = {"E", "F"}

HEADER_FONT = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2F75B5")
DATA_FONT = Font(name="Microsoft YaHei", size=10)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True)
THIN = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)


def _fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


FILL_DONE = _fill("D9EAD3")
FILL_ACTIVE = _fill("D0E4F7")
FILL_DRAFT = _fill("EFEFEF")
FILL_OVERDUE = _fill("F4CCCC")
FILL_AT_RISK = _fill("FCE5CD")
FILL_ON_TRACK = _fill("D9EAD3")
FILL_COMPLETED = _fill("B7D7A8")
FILL_UNKNOWN = _fill("EFEFEF")
FILL_PICK_FULL = _fill("E2EFDA")
FILL_PICK_PART = _fill("FFEB9C")
FILL_PICK_LOW = _fill("FFC7CE")
FILL_PICK_OVER = _fill("E4D7F5")
FILL_PICK_NONE = _fill("EFEFEF")
FILL_SO_MISS = _fill("FFF2CC")

RISK_FILL_MAP = {
    "completed": FILL_COMPLETED,
    "on_track": FILL_ON_TRACK,
    "at_risk": FILL_AT_RISK,
    "overdue": FILL_OVERDUE,
    "unknown": FILL_UNKNOWN,
}


def safe_float(value, default=0.0):
    try:
        return float(value) if value is not None else default
    except (TypeError, ValueError):
        return default


def write_header(ws, headers):
    for ci, (title, _) in enumerate(headers, 1):
        cell = ws.cell(row=1, column=ci, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN


def set_widths(ws, headers):
    for ci, (_, width) in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(ci)].width = width


def _set_cell(cell, value, fmt=None, align=LEFT, fill=None):
    cell.value = value
    cell.font = DATA_FONT
    cell.border = THIN
    cell.alignment = align
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill


def _completion_fill(is_done, raw_status):
    if is_done:
        return FILL_DONE
    if raw_status in MO_ACTIVE_STATUS:
        return FILL_ACTIVE
    return FILL_DRAFT


def _pick_fill(rate, demand):
    if demand == 0:
        return FILL_PICK_NONE
    if rate > 105:
        return FILL_PICK_OVER
    if rate >= 100:
        return FILL_PICK_FULL
    if rate >= 50:
        return FILL_PICK_PART
    return FILL_PICK_LOW


def build_bottleneck(stock_rows, relations):
    mo_to_so = {}
    mo_to_customer = {}
    for row in relations:
        mo_no = row["mo_billno"]
        so_no = row.get("so_billno", "")
        customer = row.get("so_customer_name", "")
        if so_no:
            mo_to_so.setdefault(mo_no, set()).add(so_no)
        if customer:
            mo_to_customer.setdefault(mo_no, set()).add(customer)

    agg = {}
    for stock in stock_rows or []:
        mo_no = stock.get("orderno", "") or ""
        for entry in stock.get("stockentry", []) or []:
            mat_no = entry.get("materialid_number", "") or ""
            if not mat_no:
                continue
            item = agg.setdefault(
                mat_no,
                {
                    "mat_no": mat_no,
                    "mat_name": entry.get("materialid_name", "") or "",
                    "demand": 0.0,
                    "issued": 0.0,
                    "unissued": 0.0,
                    "mos": set(),
                    "sos": set(),
                    "customers": set(),
                },
            )
            item["demand"] += safe_float(entry.get("demandqty"))
            item["issued"] += safe_float(entry.get("actissueqty"))
            item["unissued"] += safe_float(entry.get("unissueqty"))
            if mo_no:
                item["mos"].add(mo_no)
                item["sos"].update(mo_to_so.get(mo_no, set()))
                item["customers"].update(mo_to_customer.get(mo_no, set()))

    result = []
    for item in agg.values():
        if item["unissued"] <= 0:
            continue
        item["mo_count"] = len(item["mos"])
        item["so_count"] = len(item["sos"])
        item["rate"] = (item["issued"] / item["demand"] * 100) if item["demand"] > 0 else 0.0
        item["mos"] = ",".join(sorted(item["mos"]))
        item["sos"] = ",".join(sorted(item["sos"]))
        item["customers"] = ",".join(sorted(item["customers"]))
        result.append(item)
    result.sort(key=lambda item: (-item["unissued"], -item["demand"]))
    return result


def export_excel(relations, stock_rows, file_name, export_dir):
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "SO-MO关联明细"
    h1 = [
        ("序号", 6),
        ("工单号", 22),
        ("工单日期", 12),
        ("工单头状态", 16),
        ("工单行状态", 14),
        ("工单物料编码", 18),
        ("工单物料名", 30),
        ("计划数量", 12),
        ("计划开工", 16),
        ("计划完工", 16),
        ("完工状态", 20),
        ("完工率%", 10),
        ("汇报数量", 12),
        ("合格品数量", 12),
        ("已入库数量", 12),
        ("实际完工时间", 18),
        ("交货风险", 14),
        ("应领数量", 12),
        ("已领数量", 12),
        ("未领数量", 12),
        ("领料率%", 10),
        ("销售订单号", 22),
        ("SO行号", 8),
        ("SO日期", 12),
        ("SO状态", 14),
        ("SO交货状态", 14),
        ("客户", 24),
        ("SO物料编码", 16),
        ("SO数量", 12),
        ("SO含税金额", 16),
        ("SO交货日期", 12),
        ("是否关联", 10),
    ]
    write_header(ws1, h1)
    set_widths(ws1, h1)

    for row_i, row in enumerate(relations, 1):
        comp_fill = _completion_fill(row["mo_is_done"], row["mo_raw_status"])
        risk_fill = RISK_FILL_MAP.get(row["delivery_risk"], FILL_UNKNOWN)
        pick_fill = _pick_fill(row["pick_rate"], row["pick_demand"])
        so_fill = None if row["linked"] else FILL_SO_MISS
        values = [
            row_i,
            row["mo_billno"],
            row["mo_billdate"],
            row["mo_billstatus"],
            row["mo_bizstatus"],
            row["mo_material_no"],
            row["mo_material_name"],
            row["mo_qty"],
            row["mo_planbegin"],
            row["mo_planend"],
            row["mo_completion"],
            row["mo_completion_rate"],
            row["mo_rptqty"],
            row["mo_qualifiedqty"],
            row["mo_stockqty"],
            row["mo_completime"],
            row["delivery_risk_label"],
            row["pick_demand"],
            row["pick_issued"],
            row["pick_unissued"],
            round(row["pick_rate"], 1),
            row["so_billno"],
            row["so_entry_seq"],
            row["so_bizdate"],
            row["so_billstatus"],
            row["so_orderstatus"],
            row["so_customer_name"],
            row["so_material_no"],
            row["so_qty"],
            row["so_amountandtax"],
            row["so_deliverydate"],
            "Y" if row["linked"] else "N",
        ]
        for ci, value in enumerate(values, 1):
            fill = None
            if 11 <= ci <= 16:
                fill = comp_fill
            elif ci == 17:
                fill = risk_fill
            elif 18 <= ci <= 21:
                fill = pick_fill
            elif so_fill and ci > 17:
                fill = so_fill
            cell = ws1.cell(row=row_i + 1, column=ci)
            if isinstance(value, (int, float)) and ci != 1:
                fmt = "0.0" if ci in (12, 21) else "#,##0.00"
                _set_cell(cell, value, fmt=fmt, align=RIGHT, fill=fill)
            elif ci == 1:
                _set_cell(cell, value, align=CENTER, fill=fill)
            else:
                _set_cell(cell, value, align=LEFT, fill=fill)
    ws1.freeze_panes = "A2"
    ws1.auto_filter.ref = f"A1:{get_column_letter(len(h1))}1"

    ws2 = wb.create_sheet("按SO汇总")
    h2 = [
        ("销售订单号", 22),
        ("SO日期", 12),
        ("SO状态", 14),
        ("SO交货状态", 14),
        ("客户", 26),
        ("SO交货日期", 14),
        ("关联MO行数", 12),
        ("已完工行数", 12),
        ("进行中行数", 12),
        ("未下达行数", 10),
        ("MO完工率%", 12),
        ("逾期行数", 10),
        ("风险行数", 10),
        ("正常行数", 10),
        ("工单合计数量", 16),
        ("SO含税金额", 16),
    ]
    write_header(ws2, h2)
    set_widths(ws2, h2)

    so_agg = defaultdict(
        lambda: {
            "bizdate": "",
            "billstatus": "",
            "orderstatus": "",
            "customer": "",
            "delivery": "",
            "total": 0,
            "done": 0,
            "active": 0,
            "draft": 0,
            "overdue": 0,
            "at_risk": 0,
            "on_track": 0,
            "mo_qty": 0.0,
            "so_amt": 0.0,
        }
    )
    for row in relations:
        if not row["linked"]:
            continue
        item = so_agg[row["so_billno"]]
        item["bizdate"] = row["so_bizdate"]
        item["billstatus"] = row["so_billstatus"]
        item["orderstatus"] = row["so_orderstatus"]
        item["customer"] = row["so_customer_name"]
        item["delivery"] = row["so_deliverydate"]
        item["total"] += 1
        if row["mo_is_done"]:
            item["done"] += 1
        elif row["mo_raw_status"] in MO_ACTIVE_STATUS:
            item["active"] += 1
        else:
            item["draft"] += 1
        if row["delivery_risk"] == "overdue":
            item["overdue"] += 1
        elif row["delivery_risk"] == "at_risk":
            item["at_risk"] += 1
        elif row["delivery_risk"] == "on_track":
            item["on_track"] += 1
        item["mo_qty"] += row["mo_qty"]
        item["so_amt"] += row["so_amountandtax"]

    sorted_so = sorted(so_agg.items(), key=lambda item: (-item[1]["overdue"], -item[1]["at_risk"], -item[1]["total"]))
    for ri, (so_no, item) in enumerate(sorted_so, 1):
        mo_rate = (item["done"] / item["total"] * 100) if item["total"] > 0 else 0.0
        if item["overdue"] > 0:
            row_fill = FILL_OVERDUE
        elif item["at_risk"] > 0:
            row_fill = FILL_AT_RISK
        elif item["done"] == item["total"]:
            row_fill = FILL_DONE
        else:
            row_fill = None
        values = [
            so_no,
            item["bizdate"],
            item["billstatus"],
            item["orderstatus"],
            item["customer"],
            item["delivery"],
            item["total"],
            item["done"],
            item["active"],
            item["draft"],
            round(mo_rate, 1),
            item["overdue"],
            item["at_risk"],
            item["on_track"],
            item["mo_qty"],
            item["so_amt"],
        ]
        for ci, value in enumerate(values, 1):
            cell = ws2.cell(row=ri + 1, column=ci)
            fmt = "0.0" if ci == 11 else ("#,##0.00" if ci >= 15 else "#,##0")
            _set_cell(
                cell,
                value,
                fmt=fmt if isinstance(value, (int, float)) else None,
                align=RIGHT if isinstance(value, (int, float)) else LEFT,
                fill=row_fill,
            )
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:{get_column_letter(len(h2))}1"

    ws3 = wb.create_sheet("工单完工明细")
    h3 = [
        ("工单号", 22),
        ("工单日期", 12),
        ("工单头状态", 16),
        ("工单行状态", 14),
        ("工单物料", 30),
        ("计划数量", 12),
        ("完工状态", 20),
        ("完工率%", 10),
        ("汇报数量", 12),
        ("合格品", 12),
        ("已入库", 12),
        ("实际完工时间", 18),
        ("SO号", 22),
        ("SO交货日期", 14),
        ("交货风险", 14),
        ("领料率%", 10),
        ("领料状态", 12),
    ]
    write_header(ws3, h3)
    set_widths(ws3, h3)

    mo_seen = {}
    for row in relations:
        mo_seen.setdefault(row["mo_billno"], row)
    risk_order = {"overdue": 0, "at_risk": 1, "on_track": 2, "completed": 3, "unknown": 4}
    sorted_mo = sorted(mo_seen.values(), key=lambda row: (risk_order.get(row["delivery_risk"], 9), not row["mo_is_done"], row["mo_billno"]))

    for ri, row in enumerate(sorted_mo, 1):
        demand = row["pick_demand"]
        rate = row["pick_rate"]
        if demand == 0:
            pick_status = "无用料单"
        elif rate > 105:
            pick_status = "超额领"
        elif rate >= 100:
            pick_status = "齐料"
        elif rate >= 50:
            pick_status = "部分"
        else:
            pick_status = "欠料"
        values = [
            row["mo_billno"],
            row["mo_billdate"],
            row["mo_billstatus"],
            row["mo_bizstatus"],
            row["mo_material_name"],
            row["mo_qty"],
            row["mo_completion"],
            row["mo_completion_rate"],
            row["mo_rptqty"],
            row["mo_qualifiedqty"],
            row["mo_stockqty"],
            row["mo_completime"],
            row["so_billno"],
            row["so_deliverydate"],
            row["delivery_risk_label"],
            round(rate, 1),
            pick_status,
        ]
        comp_fill = _completion_fill(row["mo_is_done"], row["mo_raw_status"])
        risk_fill = RISK_FILL_MAP.get(row["delivery_risk"], FILL_UNKNOWN)
        pick_fill = _pick_fill(rate, demand)
        for ci, value in enumerate(values, 1):
            if 7 <= ci <= 12:
                fill = comp_fill
            elif ci == 15:
                fill = risk_fill
            elif ci in (16, 17):
                fill = pick_fill
            else:
                fill = None
            cell = ws3.cell(row=ri + 1, column=ci)
            fmt = "0.0" if ci in (8, 16) else ("#,##0.00" if ci in (6, 9, 10, 11) else None)
            _set_cell(cell, value, fmt=fmt if isinstance(value, (int, float)) else None, align=RIGHT if isinstance(value, (int, float)) else LEFT, fill=fill)
    ws3.freeze_panes = "A2"
    ws3.auto_filter.ref = f"A1:{get_column_letter(len(h3))}1"

    ws4 = wb.create_sheet("逾期风险工单")
    h4 = [
        ("风险等级", 14),
        ("工单号", 22),
        ("工单头状态", 14),
        ("客户", 22),
        ("SO号", 22),
        ("SO交货日期", 14),
        ("工单物料", 28),
        ("计划数量", 12),
        ("完工率%", 10),
        ("汇报数量", 12),
        ("计划完工", 16),
        ("领料率%", 10),
        ("未领数量", 12),
    ]
    write_header(ws4, h4)
    set_widths(ws4, h4)
    urgent = [row for row in sorted_mo if row["delivery_risk"] in ("overdue", "at_risk")]
    for ri, row in enumerate(urgent, 1):
        risk_fill = RISK_FILL_MAP[row["delivery_risk"]]
        values = [
            row["delivery_risk_label"],
            row["mo_billno"],
            row["mo_billstatus"],
            row["so_customer_name"],
            row["so_billno"],
            row["so_deliverydate"],
            row["mo_material_name"],
            row["mo_qty"],
            row["mo_completion_rate"],
            row["mo_rptqty"],
            row["mo_planend"],
            round(row["pick_rate"], 1),
            row["pick_unissued"],
        ]
        for ci, value in enumerate(values, 1):
            cell = ws4.cell(row=ri + 1, column=ci)
            fmt = "0.0" if ci in (9, 12) else ("#,##0.00" if ci in (8, 10, 13) else None)
            _set_cell(cell, value, fmt=fmt if isinstance(value, (int, float)) else None, align=RIGHT if isinstance(value, (int, float)) else LEFT, fill=risk_fill)
    ws4.freeze_panes = "A2"

    ws5 = wb.create_sheet("欠料明细")
    h5 = [
        ("序号", 6),
        ("工单号", 22),
        ("工单状态", 14),
        ("客户", 22),
        ("SO号", 22),
        ("用料清单号", 26),
        ("物料编码", 16),
        ("物料名称", 28),
        ("应领", 12),
        ("已领", 12),
        ("未领", 12),
        ("退料", 12),
        ("行领料率%", 10),
        ("行状态", 12),
    ]
    write_header(ws5, h5)
    set_widths(ws5, h5)
    mo_ctx = {}
    for row in relations:
        mo_ctx.setdefault(row["mo_billno"], row)
    detail_rows = []
    for stock in stock_rows or []:
        mo_no = stock.get("orderno", "") or ""
        stock_no = stock.get("billno", "") or ""
        ctx = mo_ctx.get(mo_no, {})
        for entry in stock.get("stockentry", []) or []:
            demand = safe_float(entry.get("demandqty"))
            issued = safe_float(entry.get("actissueqty"))
            unissued = safe_float(entry.get("unissueqty"))
            rejected = safe_float(entry.get("rejectedqty"))
            rate = (issued / demand * 100) if demand > 0 else 0.0
            if demand == 0 and issued == 0:
                continue
            if rate >= 100:
                continue
            if rate > 105:
                row_status = "超额"
            elif rate >= 50:
                row_status = "部分"
            elif issued > 0:
                row_status = "严重欠料"
            else:
                row_status = "未领料"
            detail_rows.append(
                {
                    "mo_no": mo_no,
                    "mo_status": ctx.get("mo_billstatus", ""),
                    "customer": ctx.get("so_customer_name", ""),
                    "so_billno": ctx.get("so_billno", ""),
                    "stock_no": stock_no,
                    "mat_no": entry.get("materialid_number", "") or "",
                    "mat_name": entry.get("materialid_name", "") or "",
                    "demand": demand,
                    "issued": issued,
                    "unissued": unissued,
                    "rejected": rejected,
                    "rate": rate,
                    "row_status": row_status,
                }
            )
    detail_rows.sort(key=lambda item: (item["rate"], item["mo_no"]))
    for ri, item in enumerate(detail_rows, 1):
        pick_fill = _pick_fill(item["rate"], item["demand"])
        values = [
            ri,
            item["mo_no"],
            item["mo_status"],
            item["customer"],
            item["so_billno"],
            item["stock_no"],
            item["mat_no"],
            item["mat_name"],
            item["demand"],
            item["issued"],
            item["unissued"],
            item["rejected"],
            round(item["rate"], 1),
            item["row_status"],
        ]
        for ci, value in enumerate(values, 1):
            cell = ws5.cell(row=ri + 1, column=ci)
            fill = pick_fill if ci in (9, 10, 11, 12, 13) else None
            fmt = "0.0" if ci == 13 else ("#,##0.00" if ci in (9, 10, 11, 12) else None)
            align = RIGHT if isinstance(value, (int, float)) else (CENTER if ci == 1 else LEFT)
            _set_cell(cell, value, fmt=fmt if isinstance(value, (int, float)) else None, align=align, fill=fill)
    ws5.freeze_panes = "A2"
    ws5.auto_filter.ref = f"A1:{get_column_letter(len(h5))}1"

    ws6 = wb.create_sheet("瓶颈物料汇总")
    h6 = [
        ("排名", 6),
        ("物料编码", 18),
        ("物料名称", 32),
        ("应领总量", 14),
        ("已领总量", 14),
        ("未领总量", 14),
        ("领料率%", 10),
        ("涉工单数", 12),
        ("涉SO数", 10),
        ("工单清单", 50),
        ("SO清单", 50),
        ("客户清单", 30),
    ]
    write_header(ws6, h6)
    set_widths(ws6, h6)
    bottleneck = build_bottleneck(stock_rows, relations)
    for ri, item in enumerate(bottleneck, 1):
        pick_fill = _pick_fill(item["rate"], item["demand"])
        values = [
            ri,
            item["mat_no"],
            item["mat_name"],
            item["demand"],
            item["issued"],
            item["unissued"],
            round(item["rate"], 1),
            item["mo_count"],
            item["so_count"],
            item["mos"],
            item["sos"],
            item["customers"],
        ]
        for ci, value in enumerate(values, 1):
            cell = ws6.cell(row=ri + 1, column=ci)
            fill = pick_fill if ci in (4, 5, 6, 7) else None
            fmt = "0.0" if ci == 7 else ("#,##0.00" if ci in (4, 5, 6) else ("0" if ci in (8, 9) else None))
            align = RIGHT if isinstance(value, (int, float)) else (CENTER if ci == 1 else LEFT)
            _set_cell(cell, value, fmt=fmt if isinstance(value, (int, float)) else None, align=align, fill=fill)
    ws6.freeze_panes = "A2"
    ws6.auto_filter.ref = f"A1:{get_column_letter(len(h6))}1"

    output_path = os.path.join(export_dir, file_name)
    wb.save(output_path)
    return output_path
